from docx import Document

# Load the existing document
doc = Document('Sleep_Detection_Project_Report.docx')

# Find the INTRODUCTION heading and replace the content after it
found_intro = False
intro_index = None

for i, para in enumerate(doc.paragraphs):
    if 'INTRODUCTION' in para.text and para.style.name == 'Heading 1':
        found_intro = True
        intro_index = i
        break

if found_intro:
    # New introduction content paragraphs
    new_intro_paras = [
        'Modern healthcare systems increasingly recognize sleep monitoring as critical for early detection of sleep disorders, cardiac arrhythmias, respiratory anomalies, and metabolic disorders. Clinical sleep assessment through polysomnography (PSG) has been the gold standard for decades, yet accessibility remains limited due to cost, infrastructure requirements, and patient discomfort. Simultaneously, the growing prevalence of sleep disorders — affecting over 50 million Americans annually — creates an urgent need for scalable, accessible, and continuous monitoring solutions.',
        
        'Sleep-related physiological disturbances — caused by obstructive sleep apnea, periodic breathing, postural changes, micro-arousals, and autonomic nervous system dysregulation — can be life-threatening if not detected early. Undiagnosed sleep apnea alone accounts for increased cardiovascular mortality, while missed REM behavior disorder signals neurodegenerative disease progression. The economic burden of undiagnosed sleep disorders exceeds $100 billion annually in lost productivity and healthcare costs.',
    ]
    
    limitations_heading = 'Limitations of Existing Methods:'
    limitations_items = [
        'Manual sleep staging requires certified sleep specialists, creating significant bottlenecks and geographic disparities in access.',
        'Contact-based sensors (electrodes, chest belts, actigraphy) cause discomfort, sleep disruption, and skin irritation, compromising measurement validity.',
        'Wearable devices provide only fragmented vital signs without comprehensive sleep stage classification and lack clinical validation.',
        'Most AI-based solutions are limited to heart rate or respiration alone, failing to capture the multi-dimensional physiological signatures of true sleep stage classification.',
    ]
    
    solution_heading = 'Proposed Solution:'
    solution_items = [
        'Detects and classifies sleep into four clinically relevant stages (Wake, REM, Light, Deep) using mmWave radar.',
        'Uses physiological-specific optimized models combining rule-based classifiers with machine learning ensemble methods.',
        'Supports real-time inference (< 50 ms per 30-second epoch).',
        'Includes a complete deployment pipeline (Flask REST API + React web dashboard).',
        'Enables continuous home-based and clinical monitoring without subject discomfort or electrode degradation.',
    ]
    
    # Remove old intro content (Background, Problem Statement, Solution Overview, Key Features sections)
    paragraphs_to_remove = []
    for i in range(intro_index + 1, min(intro_index + 50, len(doc.paragraphs))):
        para = doc.paragraphs[i]
        if para.style.name == 'Heading 1':  # Stop when we hit the next major section
            break
        paragraphs_to_remove.append(i)
    
    # Remove paragraphs in reverse order to maintain indices
    for i in reversed(paragraphs_to_remove):
        p = doc.paragraphs[i]._element
        p.getparent().remove(p)
    
    # Insert new introduction content after INTRODUCTION heading
    insert_pos = intro_index + 1
    
    # Add introduction paragraphs
    for para_text in new_intro_paras:
        new_para = doc.paragraphs[insert_pos]._element
        new_para = new_para.addprevious(doc.add_paragraph(para_text)._element)
        insert_pos += 1
    
    # Add blank line
    doc.paragraphs[insert_pos]._element.addprevious(doc.add_paragraph()._element)
    insert_pos += 1
    
    # Add Limitations heading
    limitations_head = doc.add_paragraph(limitations_heading)
    limitations_head.style = 'Normal'
    doc.paragraphs[insert_pos]._element.addprevious(limitations_head._element)
    insert_pos += 1
    
    # Add limitation bullets
    for item in limitations_items:
        bullet = doc.add_paragraph(item, style='List Bullet')
        doc.paragraphs[insert_pos]._element.addprevious(bullet._element)
        insert_pos += 1
    
    # Add blank line
    doc.paragraphs[insert_pos]._element.addprevious(doc.add_paragraph()._element)
    insert_pos += 1
    
    # Add Proposed Solution heading
    solution_head = doc.add_paragraph(solution_heading)
    solution_head.style = 'Normal'
    doc.paragraphs[insert_pos]._element.addprevious(solution_head._element)
    insert_pos += 1
    
    # Add opening line for solution
    intro_solution = doc.add_paragraph('This project introduces a novel AI-powered non-contact sleep monitoring system that:')
    doc.paragraphs[insert_pos]._element.addprevious(intro_solution._element)
    insert_pos += 1
    
    # Add solution bullets
    for item in solution_items:
        bullet = doc.add_paragraph(item, style='List Bullet')
        doc.paragraphs[insert_pos]._element.addprevious(bullet._element)
        insert_pos += 1
    
    # Save the updated document
    doc.save('Sleep_Detection_Project_Report.docx')
    print('✓ INTRODUCTION section updated successfully!')
    print('✓ New format applied with Limitations and Proposed Solution sections')
else:
    print('✗ INTRODUCTION section not found')
