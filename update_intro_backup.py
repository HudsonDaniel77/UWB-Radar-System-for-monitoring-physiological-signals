from docx import Document
import shutil
import os

# Make a copy first
shutil.copy('Sleep_Detection_Project_Report.docx', 'Sleep_Detection_Project_Report_temp.docx')

# Load the temporary copy
doc = Document('Sleep_Detection_Project_Report_temp.docx')

# Find the INTRODUCTION heading
found_intro = False
intro_index = None

for i, para in enumerate(doc.paragraphs):
    if 'INTRODUCTION' in para.text and para.style.name == 'Heading 1':
        found_intro = True
        intro_index = i
        break

if found_intro:
    print(f'Found INTRODUCTION at index {intro_index}')
    
    # Find the next major heading to know where to stop removing
    next_heading_index = None
    for i in range(intro_index + 1, len(doc.paragraphs)):
        if doc.paragraphs[i].style.name == 'Heading 1':
            next_heading_index = i
            break
    
    print(f'Next heading at index {next_heading_index}')
    
    # Remove all content between INTRODUCTION and next heading
    if next_heading_index:
        for i in range(next_heading_index - 1, intro_index, -1):
            p = doc.paragraphs[i]._element
            p.getparent().remove(p)
    
    # Now add the new content
    # Add blank line
    blank = doc.add_paragraph()
    doc.paragraphs[intro_index + 1]._element.addnext(blank._element)
    
    # Add the new paragraphs
    intro_p1 = doc.add_paragraph(
        'Modern healthcare systems increasingly recognize sleep monitoring as critical for early detection of sleep disorders, cardiac arrhythmias, respiratory anomalies, and metabolic disorders. Clinical sleep assessment through polysomnography (PSG) has been the gold standard for decades, yet accessibility remains limited due to cost, infrastructure requirements, and patient discomfort. Simultaneously, the growing prevalence of sleep disorders — affecting over 50 million Americans annually — creates an urgent need for scalable, accessible, and continuous monitoring solutions.'
    )
    doc.paragraphs[intro_index + 1]._element.addnext(intro_p1._element)
    
    intro_p2 = doc.add_paragraph(
        'Sleep-related physiological disturbances — caused by obstructive sleep apnea, periodic breathing, postural changes, micro-arousals, and autonomic nervous system dysregulation — can be life-threatening if not detected early. Undiagnosed sleep apnea alone accounts for increased cardiovascular mortality, while missed REM behavior disorder signals neurodegenerative disease progression. The economic burden of undiagnosed sleep disorders exceeds $100 billion annually in lost productivity and healthcare costs.'
    )
    doc.paragraphs[intro_index + 1]._element.addnext(intro_p2._element)
    
    # Add blank
    blank2 = doc.add_paragraph()
    doc.paragraphs[intro_index + 1]._element.addnext(blank2._element)
    
    # Add Limitations heading
    lim_heading = doc.add_paragraph('Limitations of Existing Methods:')
    doc.paragraphs[intro_index + 1]._element.addnext(lim_heading._element)
    
    # Add limitation bullets
    limitations = [
        'Manual sleep staging requires certified sleep specialists, creating significant bottlenecks and geographic disparities in access.',
        'Contact-based sensors (electrodes, chest belts, actigraphy) cause discomfort, sleep disruption, and skin irritation, compromising measurement validity.',
        'Wearable devices provide only fragmented vital signs without comprehensive sleep stage classification and lack clinical validation.',
        'Most AI-based solutions are limited to heart rate or respiration alone, failing to capture the multi-dimensional physiological signatures of true sleep stage classification.',
    ]
    
    for item in reversed(limitations):
        bullet = doc.add_paragraph(item, style='List Bullet')
        doc.paragraphs[intro_index + 1]._element.addnext(bullet._element)
    
    # Add blank
    blank3 = doc.add_paragraph()
    doc.paragraphs[intro_index + 1]._element.addnext(blank3._element)
    
    # Add Proposed Solution heading
    sol_heading = doc.add_paragraph('Proposed Solution:')
    doc.paragraphs[intro_index + 1]._element.addnext(sol_heading._element)
    
    # Add intro line
    sol_intro = doc.add_paragraph('This project introduces a novel AI-powered non-contact sleep monitoring system that:')
    doc.paragraphs[intro_index + 1]._element.addnext(sol_intro._element)
    
    # Add solution bullets
    solutions = [
        'Detects and classifies sleep into four clinically relevant stages (Wake, REM, Light, Deep) using mmWave radar.',
        'Uses physiological-specific optimized models combining rule-based classifiers with machine learning ensemble methods.',
        'Supports real-time inference (< 50 ms per 30-second epoch).',
        'Includes a complete deployment pipeline (Flask REST API + React web dashboard).',
        'Enables continuous home-based and clinical monitoring without subject discomfort or electrode degradation.',
    ]
    
    for item in reversed(solutions):
        bullet = doc.add_paragraph(item, style='List Bullet')
        doc.paragraphs[intro_index + 1]._element.addnext(bullet._element)
    
    # Save with new name
    doc.save('Sleep_Detection_Project_Report_updated.docx')
    print('✓ Updated document saved as Sleep_Detection_Project_Report_updated.docx')
    
    # Replace the original
    os.remove('Sleep_Detection_Project_Report.docx')
    os.rename('Sleep_Detection_Project_Report_updated.docx', 'Sleep_Detection_Project_Report.docx')
    os.remove('Sleep_Detection_Project_Report_temp.docx')
    print('✓ Original file updated')
    
else:
    print('✗ INTRODUCTION not found')
